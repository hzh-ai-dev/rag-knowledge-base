"""文档加载器：把各种格式的文件读成统一的 Document 列表。

统一元数据约定（后续检索、引用溯源都依赖它）：
- ``source``    文件名，用于展示与按文件删除
- ``file_path`` 绝对路径
- ``file_type`` 扩展名（不含点）
- ``doc_id``    文件内容哈希，做去重与稳定 ID
- ``page``      PDF 才有，页码从 1 开始
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from langchain_core.documents import Document

from app.utils.logging import get_logger

logger = get_logger(__name__)

# 中文文档常见编码，按顺序尝试
_TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1")


class UnsupportedFileTypeError(ValueError):
    """文件类型不支持。"""


def compute_doc_id(path: Path) -> str:
    """基于文件内容计算稳定 ID（同内容文件 ID 一致，便于增量更新）。"""
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def _read_text(path: Path) -> str:
    """按多种编码尝试解码，避免中文文档乱码。"""
    raw = path.read_bytes()
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _load_plain_text(path: Path, doc_id: str) -> list[Document]:
    text = _read_text(path).strip()
    if not text:
        return []
    return [Document(page_content=text, metadata=_base_metadata(path, doc_id))]


def _load_pdf(path: Path, doc_id: str) -> list[Document]:
    from pypdf import PdfReader

    documents: list[Document] = []
    reader = PdfReader(str(path))
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        meta = _base_metadata(path, doc_id)
        meta["page"] = index
        documents.append(Document(page_content=text, metadata=meta))
    return documents


def _load_docx(path: Path, doc_id: str) -> list[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    blocks: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        # 保留标题层级，让后续分块能感知文档结构
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            blocks.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            blocks.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    content = "\n\n".join(blocks).strip()
    if not content:
        return []
    return [Document(page_content=content, metadata=_base_metadata(path, doc_id))]


def _base_metadata(path: Path, doc_id: str) -> dict[str, object]:
    return {
        "source": path.name,
        "file_path": str(path.resolve()),
        "file_type": path.suffix.lower().lstrip("."),
        "doc_id": doc_id,
    }


_LOADERS = {
    ".txt": _load_plain_text,
    ".md": _load_plain_text,
    ".markdown": _load_plain_text,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def load_file(path: str | Path) -> list[Document]:
    """加载单个文件，返回 Document 列表（PDF 按页拆分）。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        supported = ", ".join(sorted(_LOADERS))
        raise UnsupportedFileTypeError(
            f"不支持的文件类型 {suffix}，当前支持: {supported}"
        )

    doc_id = compute_doc_id(path)
    documents = loader(path, doc_id)
    logger.info("加载 %s -> %d 个文档单元", path.name, len(documents))
    return documents


def iter_supported_files(directory: str | Path) -> list[Path]:
    """递归扫描目录下所有受支持的文件。"""
    directory = Path(directory)
    if not directory.exists():
        return []
    files: list[Path] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in _LOADERS:
            files.append(path)
    return files
